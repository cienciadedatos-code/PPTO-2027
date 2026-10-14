"""Paso 7. Hojas de gerentes: cantidad y precio 2027 prellenados por cliente y descripción.

Entrada : consulta/cache/ventas.parquet, proyeccion_base.parquet y consulta/catalogo.xlsx
Salida  : output/revision_gerentes/Revision_<LINEA>.xlsx   (una por línea; se respeta lo que escriban)
          consulta/cache/filas_gerentes.parquet            (lo revisado, para los pasos siguientes)
Correr  : py src/07_gerentes.py   (una vez para generar las hojas y otra vez cuando regresen llenas)
"""
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

import comun as c

PARTE_A = 0.80              # las filas que acumulan este % de la venta de la línea son prioridad A
CAMBIO_FUERTE = 30          # % de cambio en cantidad o precio que se lista para confirmar
FILAS_EN_BLANCO = 40        # renglones vacíos al final de cada hoja para dar de alta clientes o productos

cfg = c.cargar_config()
CARPETA = cfg["carpeta_output"] / "revision_gerentes"
CARPETA.mkdir(exist_ok=True)
c.titulo("PASO 7. HOJAS DE GERENTES: cantidad y precio 2027 por cliente y descripción")
ventas = c.leer_cache(cfg, "ventas", "01_ventas.py")
p = c.leer_cache(cfg, "proyeccion_base", "06_proyeccion_2027.py")
cat = pd.read_excel(cfg["carpeta_consulta"] / "catalogo.xlsx", dtype=str).fillna("")
v, ult, FIN = c.mezcla_reciente(ventas, cat)
LLAVE = ["linea", "cliente_id", "producto"]

# ------------------------------------------------------------------ 1. Repartir cada clase entre cliente y descripción
mix = ult.groupby(["linea", "clase_et", "cliente_id", "producto"], as_index=False).agg(
    cliente=("cliente", "last"), unidad=("unidad", "last"), key_ref=("key_id", "last"),
    venta_12m=("venta", "sum"), cantidad_12m=("cantidad", "sum"))
mix = mix[mix["venta_12m"] > 0]
mix["parte"] = mix["venta_12m"] / mix.groupby(["linea", "clase_et"])["venta_12m"].transform("sum")
d = mix.merge(p[["linea", "clase_et", "grupo", "base_2027"]], on=["linea", "clase_et"], how="inner")
d["venta_sugerida"] = d["base_2027"] * d["parte"]

# Historia de cada fila: cantidad y venta por año, y meses con compra en los últimos 12
for anio in (2024, 2025):
    h = v[v["anio"] == anio].groupby(LLAVE).agg(**{f"cantidad_{anio}": ("cantidad", "sum"), f"venta_{anio}": ("venta", "sum")})
    d = d.merge(h, on=LLAVE, how="left")
compras = ult[ult["venta"] > 0].assign(per=lambda x: x["fecha"].dt.to_period("M")).groupby(LLAVE)["per"].nunique()
d["meses_con_compra"] = [int(compras.get(k, 0)) for k in zip(d["linea"], d["cliente_id"], d["producto"])]
d["confianza"] = pd.cut(d["meses_con_compra"], [-1, 2, 5, 12], labels=["BAJA", "MEDIA", "ALTA"]).astype(str)

# Precio sugerido: el promedio de los últimos 6 meses de ese cliente con ese producto;
# si no compró en 6 meses, el de 12; si tampoco, su última factura; y al final el promedio del producto.
fac = v[(v["tipo_doc"] == "Factura") & (v["cantidad"] > 0) & (v["venta"] > 0)].sort_values("fecha")


def precio_medio(meses):
    t = fac[fac["fecha"] > FIN - pd.DateOffset(months=meses)].groupby(LLAVE).agg(v=("venta", "sum"), q=("cantidad", "sum"))
    return (t["v"] / t["q"]).rename(f"p{meses}")


d = d.merge(precio_medio(6), on=LLAVE, how="left").merge(precio_medio(12), on=LLAVE, how="left")
ultima = fac.groupby(LLAVE).agg(p_ult=("precio_usd", "last"))
prod = fac[fac["fecha"] > FIN - pd.DateOffset(months=12)].groupby("producto").agg(v=("venta", "sum"), q=("cantidad", "sum"))
d = d.merge(ultima, on=LLAVE, how="left").merge((prod["v"] / prod["q"]).rename("p_prod"), on="producto", how="left")
d["precio_12m"] = d["p12"]
d["precio_sugerido"] = d["p6"].fillna(d["p12"]).fillna(d["p_ult"]).fillna(d["p_prod"])
d["cantidad_sugerida"] = d["venta_sugerida"] / d["precio_sugerido"]

# Clases con proyección pero sin ventas recientes que repartir (la reserva de cada línea): van en dólares
con_filas = set(zip(d["linea"], d["clase_et"]))
sin = p[(p["base_2027"] > 0) & ~pd.Series([k in con_filas for k in zip(p["linea"], p["clase_et"])], index=p.index)]
extra = pd.DataFrame({"linea": sin["linea"], "clase_et": sin["clase_et"], "grupo": sin["grupo"], "cliente_id": "POR DEFINIR",
                      "cliente": "POR DEFINIR", "producto": "POR DEFINIR", "unidad": "USD", "key_ref": "",
                      "venta_sugerida": sin["base_2027"], "cantidad_sugerida": sin["base_2027"], "precio_sugerido": 1.0,
                      "confianza": "RESERVA", "meses_con_compra": 0})
d = pd.concat([d, extra], ignore_index=True)
sin_precio = d["precio_sugerido"].isna()
d.loc[sin_precio, ["cantidad_sugerida", "precio_sugerido", "unidad"]] = [np.nan, 1.0, "USD"]
d.loc[sin_precio, "cantidad_sugerida"] = d.loc[sin_precio, "venta_sugerida"]
d = d[d["venta_sugerida"] > 0.5].copy()
orden = d.sort_values("venta_sugerida", ascending=False)
antes = orden.groupby("linea")["venta_sugerida"].cumsum() - orden["venta_sugerida"]
d["prioridad"] = (antes < PARTE_A * orden.groupby("linea")["venta_sugerida"].transform("sum")).map({True: "A", False: "B"})
d["id_fila"] = d["linea"] + "|" + d["cliente_id"] + "|" + d["clase_et"] + "|" + d["producto"]
d["cantidad_2027"], d["precio_2027"], d["motivo"], d["alta"] = d["cantidad_sugerida"], d["precio_sugerido"], "", False

COLS = ["prioridad", "confianza", "cliente", "clase_et", "producto", "unidad", "cantidad_2024", "venta_2024",
        "cantidad_2025", "venta_2025", "cantidad_12m", "venta_12m", "precio_12m", "cantidad_sugerida",
        "precio_sugerido", "cantidad_2027", "precio_2027", "venta_2027", "motivo", "id_fila"]
EDITABLES = ["cantidad_2027", "precio_2027", "motivo"]
resumen_clase = p[p["base_2027"] > 0][["linea", "clase_et", "grupo", "venta_2025", "ultimos_12m", "cierre_2026",
                                       "base_2027", "bajo_2027", "alto_2027", "clientes_12m", "cliente_principal"]]


def igual(a, b):
    return (a - b).abs() <= 0.005 * b.abs() + 1e-9


# ------------------------------------------------------------------ 2. Una hoja por línea (respetando lo ya escrito)
c.subtitulo("Hojas por línea")
altas_malas = []
for linea in c.LINEAS:
    corto = linea.replace("BU-", "").replace(" ", "_")
    ruta = CARPETA / f"Revision_{corto}.xlsx"
    g = d[d["linea"] == linea].copy()
    nuevas = pd.DataFrame()
    if ruta.exists():
        prev = pd.read_excel(ruta, sheet_name="filas")
        con_id = prev[prev["id_fila"].notna()].drop_duplicates("id_fila").set_index("id_fila")
        for col, sug in (("cantidad_2027", "cantidad_sugerida"), ("precio_2027", "precio_sugerido")):
            escrito = pd.to_numeric(g["id_fila"].map(con_id[col]), errors="coerce")
            sugerido_antes = pd.to_numeric(g["id_fila"].map(con_id[sug]), errors="coerce")
            tocado = escrito.notna() & ~igual(escrito, sugerido_antes)       # el gerente cambió ese número
            g[col] = escrito.where(tocado, g[col])
        g["motivo"] = g["id_fila"].map(con_id["motivo"]).fillna("").astype(str).replace("nan", "")
        # Altas: renglones agregados al final, sin id_fila
        al = prev[prev["id_fila"].isna() & prev["cliente"].notna()].copy()
        for col in ("cantidad_2027", "precio_2027"):
            al[col] = pd.to_numeric(al[col], errors="coerce")
        mala = al[["clase_et", "producto", "cantidad_2027", "precio_2027"]].isna().any(axis=1)
        altas_malas.append(al[mala].assign(archivo=ruta.name))
        al = al[~mala]
        if len(al):
            nuevas = pd.DataFrame({
                "linea": linea, "clase_et": al["clase_et"].map(c.norm), "grupo": "ALTA", "cliente_id": "ALTA-" + al["cliente"].map(c.norm),
                "cliente": al["cliente"].astype(str), "producto": al["producto"].map(c.norm), "unidad": al["unidad"].fillna("").astype(str),
                "key_ref": "", "venta_sugerida": 0.0, "cantidad_sugerida": 0.0, "precio_sugerido": al["precio_2027"],
                "confianza": "ALTA DEL GERENTE", "meses_con_compra": 0, "prioridad": "A", "id_fila": np.nan,
                "cantidad_2027": al["cantidad_2027"], "precio_2027": al["precio_2027"],
                "motivo": al["motivo"].fillna("").astype(str), "alta": True})
    g = g.sort_values(["prioridad", "venta_sugerida"], ascending=[True, False])
    hoja = pd.concat([g, nuevas], ignore_index=True) if len(nuevas) else g
    hoja = hoja.assign(venta_2027=np.nan)[COLS]
    pend = altas_malas[-1].reindex(columns=COLS) if altas_malas and len(altas_malas[-1]) else None
    if pend is not None:
        hoja = pd.concat([hoja, pend], ignore_index=True)
    c.escribir_excel({"filas": hoja, "por_clase": resumen_clase[resumen_clase["linea"] == linea].sort_values("base_2027", ascending=False)},
                     ruta, anchos={"cliente": 34, "producto": 52, "clase_et": 28, "motivo": 40, "id_fila": 10, "cliente_principal": 32})
    # Fórmula de venta (cantidad por precio) y color en las columnas que llena el gerente
    wb = load_workbook(ruta)
    ws = wb["filas"]
    letra = {col: get_column_letter(i + 1) for i, col in enumerate(COLS)}
    q, pr, vt = letra["cantidad_2027"], letra["precio_2027"], letra["venta_2027"]
    for r in range(2, len(hoja) + 2 + FILAS_EN_BLANCO):
        ws[f"{vt}{r}"] = f'=IF(OR({q}{r}="",{pr}{r}=""),"",{q}{r}*{pr}{r})'
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for col in EDITABLES:
        for r in range(1, len(hoja) + 2 + FILAS_EN_BLANCO):
            ws[f"{letra[col]}{r}"].fill = amarillo
    wb.save(ruta)
    d = pd.concat([d[d["linea"] != linea], g, nuevas], ignore_index=True)
    print(f"{ruta.name}: {len(g):,} filas ({int((g['prioridad'] == 'A').sum())} de prioridad A), {len(nuevas)} altas")

malas = pd.concat(altas_malas, ignore_index=True) if altas_malas else pd.DataFrame()
if len(malas):
    c.aviso(f"{len(malas)} altas incompletas (no entran). Necesitan cliente, clase_et, producto, cantidad_2027 y precio_2027:")
    c.mostrar(malas[["archivo", "cliente", "clase_et", "producto", "cantidad_2027", "precio_2027"]])

# ------------------------------------------------------------------ 3. Resultado de la revisión
d["venta_final"] = d["cantidad_2027"] * d["precio_2027"]
d["cambio_cantidad_%"] = (100 * (d["cantidad_2027"] / d["cantidad_sugerida"].replace(0, np.nan) - 1)).round(1)
d["cambio_precio_%"] = (100 * (d["precio_2027"] / d["precio_sugerido"].replace(0, np.nan) - 1)).round(1)
d["ajustada"] = d["alta"] | (d["cambio_cantidad_%"].abs() > 0.5) | (d["cambio_precio_%"].abs() > 0.5)
c.guardar_cache(cfg, "filas_gerentes", d[["linea", "clase_et", "grupo", "cliente_id", "cliente", "producto", "unidad", "key_ref",
                                           "prioridad", "confianza", "cantidad_sugerida", "precio_sugerido", "venta_sugerida",
                                           "cantidad_2027", "precio_2027", "venta_final", "ajustada", "alta", "motivo"]])

c.subtitulo("Cuántas filas y de qué tipo")
t = d[~d["alta"]].groupby(["prioridad", "confianza"]).agg(filas=("venta_sugerida", "size"), venta_sugerida=("venta_sugerida", "sum")).reset_index()
t["% de la venta"] = (100 * t["venta_sugerida"] / t["venta_sugerida"].sum()).round(1)
c.mostrar(t)
print("Prioridad A = las filas que hacen el 80% de la venta de su línea. Confianza = qué tan seguido compra el cliente")
print("(ALTA: 6 o más meses del último año; MEDIA: 3 a 5; BAJA: 1 o 2).")

c.subtitulo("Cambios de los gerentes")
aj = d[d["ajustada"]]
if len(aj) == 0:
    print("Todavía no hay cambios. Comparte las hojas de output/revision_gerentes. En las columnas amarillas cada gerente")
    print("corrige 'cantidad_2027' y 'precio_2027' donde no esté de acuerdo y escribe el 'motivo'. La venta se calcula sola.")
else:
    c.mostrar(aj.sort_values("venta_final", ascending=False)[["linea", "cliente", "producto", "cantidad_sugerida", "cantidad_2027",
              "cambio_cantidad_%", "precio_sugerido", "precio_2027", "cambio_precio_%", "venta_final", "motivo"]], 40)
    sin_motivo = aj[aj["motivo"].str.strip() == ""]
    (c.ok if len(sin_motivo) == 0 else c.aviso)(f"Cambios sin motivo: {len(sin_motivo)}. Pide al gerente que lo escriba")
    fuertes = aj[~aj["alta"] & ((aj["cambio_cantidad_%"].abs() >= CAMBIO_FUERTE) | (aj["cambio_precio_%"].abs() >= CAMBIO_FUERTE))]
    (c.ok if len(fuertes) == 0 else c.aviso)(f"Cambios de {CAMBIO_FUERTE}% o más en cantidad o precio: {len(fuertes)}. "
                                             "No es error, pero confirma que el motivo lo explique")

c.subtitulo("Proyección 2027 por línea (millones USD)")
res = d.groupby("linea").agg(base_2027=("venta_sugerida", "sum"), gerentes_2027=("venta_final", "sum"),
                             filas_cambiadas=("ajustada", "sum")).reindex(c.LINEAS)
res.loc["TOTAL"] = res.sum()
vista = (res[["base_2027", "gerentes_2027"]] / 1e6).round(2)
vista["filas_cambiadas"] = res["filas_cambiadas"].astype(int)
vista["gerentes vs base %"] = (100 * (res["gerentes_2027"] / res["base_2027"] - 1)).round(1)
print(vista.to_string())

c.subtitulo("Controles")
esperado, obtenido = p["base_2027"].sum(), d.loc[~d["alta"], "venta_sugerida"].sum()
(c.ok if abs(esperado - obtenido) < 100 else c.aviso)(f"Las filas suman la proyección base: {obtenido:,.0f} de {esperado:,.0f}")
neg = int(((d["cantidad_2027"] < 0) | (d["precio_2027"] < 0)).sum())
(c.ok if neg == 0 else c.aviso)(f"Filas con cantidad o precio negativo: {neg}")
print(f"\nHojas en: {CARPETA}")
c.cierre("PASO 7")
